import json
import os
import smtplib
import sys
import datetime
import qdarkstyle
import requests
import dotenv
import alternative
import helpers
from PyQt5 import QtWidgets, QtCore, QtGui
from PyQt5.QtWidgets import QFileDialog, QMessageBox, QTableWidgetItem, QListWidgetItem, QDialog, QDialogButtonBox, \
    QComboBox
from interface.entry_ui import Ui_Enter
from interface.window_ui import Ui_MainWindow
from interface.user_ui import Ui_Newuser
from interface.hardware_ui import Ui_Newhard
from interface.alternative_ui import Ui_Alternative
from interface.request_ui import Ui_NewRequest
from interface.email_ui import Ui_Email
from main import start
from xlwt import Workbook
from docx import Document

PATH = os.getcwd()
DB_ACCESS_TOKEN = "Basic NVJOWUJkTGR1VER4UUNjTThZWXJiNW5BOkg0ZFNjQXlHYlM4OUtnTGdaQnMydlBzaw=="
DB_URL = "https://helow19274.ru/aip/api"

def format_date(date: str) -> str:
    """
    Преобразует строку с датой формата ДД.ММ.ГГГГ в ISO-формат с текущим временем.

    :param date: Дата в формате "ДД.ММ.ГГГГ"
    :return: Строка с датой в ISO-формате.
    """
    now = datetime.datetime.now()
    day, month, year = map(int, date.split('.'))
    new_date = datetime.datetime(year, month, day, now.hour, now.minute, now.second, now.microsecond)
    return new_date.isoformat()

def set_keys(email: str, password: str) -> None:
    """
    Записывает учетные данные почты в файл .env.
    
    :param email: Логин почты
    :param password: Пароль почты
    """
    dotenv.set_key(f'{PATH}/.env', "EMAIL_USERNAME", email)
    dotenv.set_key(f'{PATH}/.env', "EMAIL_PASSWORD", password)

def check_email_credentials(email_username: str, email_password: str) -> bool:
    """
    Проверяет учетные данные почты через SMTP.
    
    :param email_username: Логин почты
    :param email_password: Пароль почты
    :return: True, если данные верны, иначе False.
    """
    smtp_server = "smtp.yandex.ru"
    try:
        mail = smtplib.SMTP_SSL(smtp_server)
        mail.login(email_username, email_password)
        mail.quit()
        return True
    except smtplib.SMTPAuthenticationError:
        return False

class LoginWindow(QtWidgets.QMainWindow):
    """
    Главное окно входа в систему.
    """
    def __init__(self):
        """
        Инициализация окна входа.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/entry_icon.png'))
        self.table_window = Ui_MainWindow
        self.ui = Ui_Enter()
        self.ui.setupUi(self)
        self.ui.btn_entry.clicked.connect(self.check)
        self.ui.password_line.returnPressed.connect(self.check)
        self.email_username = os.getenv("EMAIL_USERNAME")
        self.ui.entry_last_button.setToolTip(self.email_username)
        self.ui.entry_last_button.clicked.connect(self.previous_session)

    def previous_session(self):
        """
        Вход в предыдущую сессию, если она существует.
        """
        if self.email_username:
            dotenv.load_dotenv(f'{PATH}/.env')
            email_username = os.getenv("EMAIL_USERNAME")
            email_password = os.getenv("EMAIL_PASSWORD")
            check_email_credentials(email_username, email_password)
            self.close()
            self.table_window = MainWindow()
            self.table_window.user_table()
            hardware = helpers.get_request('hardware')
            boards = [x['name'] for x in hardware]
            self.table_window.ui.hardware_list.addItems(boards)
            self.table_window.show()
        else:
            QMessageBox.warning(
                self,
                "Ошибка",
                "Прошлой сессии не существует",
                QMessageBox.StandardButton.Ok
            )

    def check(self):
        """
        Проверка учетных данных почты и вход в систему.
        """
        email = self.ui.email_line.text()
        password = self.ui.password_line.text()
        if check_email_credentials(email, password):
            set_keys(email, password)
            self.close()
            self.table_window = MainWindow()
            self.table_window.user_table()
            hardware = helpers.get_request('hardware')
            boards = [x['name'] for x in hardware]
            self.table_window.ui.hardware_list.addItems(boards)
            self.table_window.show()
        else:
            QMessageBox.warning(
                self,
                "Ошибка",
                "Неверная почта или пароль",
                QMessageBox.StandardButton.Ok
            )


class HardDialog(QDialog):
    """
    Диалоговое окно для добавления новой аппаратной платы.
    """
    def __init__(self):
        """
        Инициализация интерфейса диалогового окна.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/hardware_icon.png'))
        self.ui = Ui_Newhard()
        self.ui.setupUi(self)
        self.setWindowFlag(QtCore.Qt.WindowContextHelpButtonHint, False)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)
        self.ui.spec_radio.pressed.connect(self.add_spec)

    def add_spec(self):
        """
        Отображает или скрывает дополнительные параметры спецификации.
        """
        if self.ui.pin_label.isHidden():
            elements = [
                (self.ui.pin_label, 7), (self.ui.pin_line, 8),
                (self.ui.log_label, 9), (self.ui.log_line, 10),
                (self.ui.mult_label, 11), (self.ui.mult_line, 12),
                (self.ui.pll_label, 13), (self.ui.pll_line, 14),
                (self.ui.memory_label, 15), (self.ui.memory_line, 16)
            ]
            for widget, row in elements:
                self.ui.gridLayout.addWidget(widget, row, 0, 1, 2)
                widget.show()
            box = self.ui.buttonBox
            self.ui.gridLayout.removeWidget(self.ui.buttonBox)
            self.ui.gridLayout.addWidget(box, 17, 1, 1, 1)
        else:
            for widget in [
                self.ui.memory_line, self.ui.memory_label, self.ui.pll_line, self.ui.pll_label,
                self.ui.mult_line, self.ui.mult_label, self.ui.pin_label, self.ui.pin_line,
                self.ui.log_label, self.ui.log_line
            ]:
                widget.hide()
            self.adjustSize()

    def check(self):
        """
        Проверяет введенные данные и отправляет запрос на создание новой аппаратной платы.
        """
        specs = {
            "log_elems": self.ui.log_line.text(),
            "memory": self.ui.memory_line.text(),
            "pll": self.ui.pll_line.text(),
            "multiplier": self.ui.mult_line.text(),
            "pins": self.ui.pin_line.text()
        } if self.ui.spec_radio.isChecked() else {}
        
        hardware_data = {
            "name": self.ui.name_line.text(),
            "type": self.ui.type_line.currentText(),
            "description": self.ui.description_line.text(),
            "image_link": "https://encrypted-tbn0.gstatic.com/images?q=tbn:ANd9GcStW9UvyhB2beq-tiyJMhzWdP98Rny7PzRaPA&usqp=CAU",
            "specifications": specs
        }

        hardware_response = helpers.post_request('hardware', hardware_data)
        if 'detail' in hardware_response:
            error_messages = {
                'type': "Ошибка в вводе типа платы",
                'name': "Ошибка в вводе названия",
                'description': "Ошибка в вводе описания. Если оно отсутствует, введите -",
                'specifications': "Ошибка в вводе спецификации"
            }
            for key, message in error_messages.items():
                if key in hardware_response['detail'][0]['loc']:
                    return QMessageBox.warning(self, "Ошибка", message, QMessageBox.StandardButton.Ok)
            return QMessageBox.warning(self, 'Неизвестная ошибка', str(hardware_response))
        else:
            self.close()


class UserDialog(QDialog):
    """
    Диалоговое окно для добавления нового пользователя.
    """
    def __init__(self):
        """
        Инициализация интерфейса диалогового окна.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/user_icon.png'))
        self.ui = Ui_Newuser()
        self.ui.setupUi(self)
        self.setWindowFlag(QtCore.Qt.WindowContextHelpButtonHint, False)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        """
        Проверяет введенные данные и отправляет запрос на добавление нового пользователя.
        """
        user_data = {
            "active": True,
            "type": self.ui.access_box.currentText(),
            "first_name": self.ui.name_line.text(),
            "last_name": self.ui.surname_line.text(),
            "patronymic": self.ui.patr_line.text(),
            "image_link": "https://clck.ru/34TfSF",
            "email": self.ui.email_line.text(),
            "phone": self.ui.phone_line.text(),
            "card_id": "string",
            "card_key": "string",
            "comment": self.ui.group_line.text()
        }

        response = helpers.post_request('user', user_data)
        if 'detail' in response:
            error_messages = {
                'email': "Ошибка в вводе почты",
                'phone': "Ошибка в вводе телефона"
            }
            for key, message in error_messages.items():
                if key in response['detail'][0]['type']:
                    return QMessageBox.warning(self, "Ошибка", message, QMessageBox.StandardButton.Ok)
            print("Еще какая-то ошибка")
            print(response)
        else:
            self.close()


class AlternativeDialog(QDialog):
    """
    Диалоговое окно для задания параметров альтернативных плат.
    """
    def __init__(self):
        """
        Инициализация интерфейса диалогового окна.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/alter_icon.png'))
        self.ui = Ui_Alternative()
        self.ui.setupUi(self)
        self.setWindowFlag(QtCore.Qt.WindowContextHelpButtonHint, False)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.make_json)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def make_json(self):
        """
        Собирает данные из полей ввода и сохраняет их в JSON-файл.
        """
        log_value = int(self.ui.log_spin.text())
        log_index = float(self.ui.log_spin2.text().replace(',', '.'))
        mem_value = int(self.ui.mem_spin.text())
        mem_index = float(self.ui.mem_spin2.text().replace(',', '.'))
        mult_value = int(self.ui.mult_spin.text())
        mult_index = float(self.ui.mult_spin2.text().replace(',', '.'))
        pll_value = int(self.ui.pll_spin.text())
        pll_index = float(self.ui.pll_spin2.text().replace(',', '.'))
        pin_value = int(self.ui.pin_spin.text())
        pin_index = float(self.ui.pin_spin2.text().replace(',', '.'))
        data = {
            "log_elems": [log_value, log_index],
            "memory": [mem_value, mem_index],
            "pll": [pll_value, pll_index],
            "multiplier": [mult_value, mult_index],
            "pins": [pin_value, pin_index]
        }
        with open(f'{PATH}/alternative/max_variance.json', 'w', encoding='utf-8') as file:
            json.dump(data, file, indent=4, ensure_ascii=False)


class RequestDialog(QDialog):
    """
    Диалоговое окно для создания запроса на оборудование.
    """
    def __init__(self):
        """
        Инициализация интерфейса диалогового окна.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/request_icon.png'))
        self.ui = Ui_NewRequest()
        self.ui.setupUi(self)
        self.setWindowFlag(QtCore.Qt.WindowContextHelpButtonHint, False)
        self.ui.buttonBox.button(QDialogButtonBox.Ok).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        """
        Проверяет введенные данные и отправляет запрос на создание нового запроса.
        """
        name, last_name = self.ui.name_line.text(), self.ui.last_name_line.text()
        email, phone = self.ui.email_line.text(), self.ui.phone_line.text()
        response = helpers.get_request('user')
        user_id = next((x['id'] for x in response if x.get('first_name') == name and x.get('last_name') == last_name), 0)
        
        if user_id == 0:
            user_data = {
                "active": True, "type": 'user', "first_name": name, "last_name": last_name,
                "patronymic": '', "image_link": "https://clck.ru/34TfSF", "email": email,
                "phone": phone, "card_id": "string", "card_key": "string", "comment": ''
            }
            user_response = helpers.post_request('user', user_data)
            if 'detail' in user_response:
                error_messages = {'email': "Ошибка в вводе почты", 'phone': "Ошибка в вводе телефона"}
                for key, message in error_messages.items():
                    if key in user_response['detail'][0]['type']:
                        return QMessageBox.warning(self, "Ошибка", message, QMessageBox.StandardButton.Ok)
                return QMessageBox.warning(self, 'Неизвестная ошибка', str(user_response))
            user_id = user_response['id']
        
        if user_id:
            locations = {x['name']: x['id'] for x in helpers.get_request('location')}
            cabinets = locations.get(self.ui.cab_box.currentText())
            
            hardware_list = {x['name']: x['id'] for x in helpers.get_request('hardware')}
            hardware = hardware_list.get(self.ui.hardware_box.currentText())
            
            data = {
                "user": user_id, "location": cabinets, "status": "new", "comment": self.ui.comment_line.text(),
                "taken_date": format_date(self.ui.date1_edit.text()),
                "return_date": format_date(self.ui.date2_edit.text()), "issued_by": user_id,
                "hardware": [{"hardware": hardware, "count": self.ui.count_box.text()}]
            }
            response = helpers.post_request('request', data)
            print(response)


class Email(QDialog):
    """
    Диалоговое окно для изменения почты.
    """
    def __init__(self):
        """
        Инициализация интерфейса диалогового окна.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/Logo_Email.jpg'))
        self.table_window = Ui_MainWindow
        self.ui = Ui_Email()
        self.ui.setupUi(self)
        self.setWindowFlag(QtCore.Qt.WindowContextHelpButtonHint, False)
        self.ui.buttonBox.button(QDialogButtonBox.Save).clicked.connect(self.check)
        self.ui.buttonBox.button(QDialogButtonBox.Cancel).clicked.connect(self.close)

    def check(self):
        """Проверка почты"""
        email = self.ui.email_line.text()
        password = self.ui.password_line.text()
        if check_email_credentials(email, password):
            set_keys(email, password)
        else:
            QMessageBox.warning(
                self,
                "Ошибка",
                "Неверная почта или пароль",
                QMessageBox.StandardButton.Ok
            )


class MainWindow(QtWidgets.QMainWindow):
    """
    Главное окно приложения.
    """
    def __init__(self):
        """
        Инициализация главного окна и подключение обработчиков событий.
        """
        super().__init__()
        self.setWindowIcon(QtGui.QIcon('interface/icons/main_icon.jpg'))
        self.email_dialog = Ui_Email
        self.status_box = QComboBox()
        self.request_dialog = Ui_NewRequest
        self.alter_dialog = Ui_Alternative
        self.hard_dialog = Ui_Newhard
        self.user_dialog = Ui_Newuser
        self.ui = Ui_MainWindow()
        self.ui.setupUi(self)
        self.ui.tableWidget.horizontalHeader().setSectionResizeMode(QtWidgets.QHeaderView.ResizeMode.Stretch)
        self.ui.tableWidget.verticalHeader().setSectionResizeMode(QtWidgets.QHeaderView.ResizeMode.Stretch)
        self.ui.action_excel.triggered.connect(self.save_excel_file)
        self.ui.action_doc.triggered.connect(self.save_doc_file)
        self.ui.action_theme.triggered.connect(self.switch_theme)
        self.ui.action_alternative.triggered.connect(self.set_alternatives)
        self.ui.action_email.triggered.connect(self.configure_email)

        self.ui.users_button.clicked.connect(self.user_table)
        self.ui.hardware_button.clicked.connect(self.hardware_table)
        self.ui.request_button.clicked.connect(self.request_table)
        self.ui.check_email_button.clicked.connect(self.accounting)
        self.ui.list_button.clicked.connect(self.check_availability)

        self.ui.add_user_button.clicked.connect(self.add_user)
        self.ui.add_hardware_button.clicked.connect(self.add_hardware)
        self.ui.add_request_button.clicked.connect(self.add_request)

        self.ui.search_button.clicked.connect(self.search_user)
        self.ui.del_button.clicked.connect(self.delete)
        self.ui.tableWidget.itemChanged.connect(self.save_edits)
        self.boxes = {}

    def accounting(self):
        """
        Проверка сообщений на почте
        """
        message = start()
        return QMessageBox.warning(
            self,
            "Результат проверки почты",
            message,
            QMessageBox.StandardButton.Ok
        )

    def check_availability(self):
        """
        Проверка наличия платы на складе
        """
        current_item = self.ui.hardware_list.currentItem()
        if current_item is None:
            return QMessageBox.warning(self, 'Внимание', 'Пожалуйста выберете плату для проверки')
        hardware_name = current_item.text()
        hardwares = helpers.get_request('hardware')
        stock = helpers.get_request('stocks')
        hw_id = None
        current_hardware = {}
        for hw in hardwares:
            if hw.get('name') == hardware_name:
                hw_id = hw.get('id')
                current_hardware = hw
        if hw_id is None:
            return QMessageBox.warning(self, 'Ошибка', 'Такой платы не существует')
        count = 0
        for st in stock:
            st_id = st['hardware']
            if st_id == hw_id:
                count += st['count']
        if count == 0:
            alter = alternative.find_alternative_board(current_hardware, hardwares)
            for hw in hardwares:
                if hw.get('name') == alter:
                    hw_id = hw.get('id')
            count = 0
            for st in stock:
                st_id = st['hardware']
                if st_id == hw_id:
                    count += st['count']

            message = QMessageBox()
            message.setWindowTitle("Наличие платы")
            if alter != '' and count != 0:
                message.setText(
                    f'Данной платы нет в наличии, но доступна альтернатива:\n{alter} в количестве {count} штук')
            else:
                message.setText('Данной платы нет в наличии, а также нет доступных альтернатив')
            message.setIcon(QMessageBox.Icon.Information)
            message.setStandardButtons(QMessageBox.Ok)
            message.exec_()

        else:
            message = QMessageBox()
            message.setWindowTitle("Наличие платы")
            message.setText(f'Доступно {count} плат {hardware_name}')
            message.setIcon(QMessageBox.Icon.Information)
            message.setStandardButtons(QMessageBox.Ok)
            message.exec_()

    def configure_email(self):
        """
        Открытие окна настройки электронной почты.
        """
        self.email_dialog = Email()
        self.email_dialog.show()

    def add_request(self):
        """
        Открытие окна добавления нового запроса.
        """
        self.request_dialog = RequestDialog()
        hardware = helpers.get_request('hardware')
        boards = [x['name'] for x in hardware]
        self.request_dialog.ui.hardware_box.addItems(boards)
        locations = helpers.get_request('location')
        locs = [x['name'] for x in locations]
        self.request_dialog.ui.cab_box.addItems(locs)
        self.request_dialog.show()

    def set_alternatives(self):
        """
        Открытие окна задания альтернативных параметров.
        """
        self.alter_dialog = AlternativeDialog()
        self.alter_dialog.show()

    def save_edits(self):
        """
        Сохранение в базе данных произведенных изменений в приложении
        """
        if self.ui.tableWidget.currentItem() is not None:
            item = self.ui.tableWidget.currentItem().text()
            current_row = self.ui.tableWidget.currentRow()
            current_col = self.ui.tableWidget.currentColumn()
            update_id = int(self.ui.tableWidget.item(current_row, 0).text())
            if self.ui.tableWidget.columnCount() == 8:
                user_dict = {"Имя": 'first_name', "Фамилия": 'last_name', "Отчество": 'patronymic',
                             "Группа": 'comment', "Уровень доступа": 'type', "Телефон": 'phone', "Почта": 'email'}

                col_name = user_dict[self.ui.tableWidget.horizontalHeaderItem(current_col).text()]
                user_response = helpers.patch_request('user', col_name, item, update_id)
                if user_response['detail'] == "OK":
                    QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
                else:
                    QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка', QMessageBox.Ok)
                    self.user_table()
            elif self.ui.tableWidget.columnCount() == 9:
                hardware_dict = {"Тип": 'type', "Название": 'name', "Описание": 'description'}

                if (self.ui.tableWidget.horizontalHeaderItem(current_col).text()) in hardware_dict:
                    col_name = hardware_dict[self.ui.tableWidget.horizontalHeaderItem(current_col).text()]
                    hardware_response = helpers.patch_request('hardware', col_name, item, update_id)
                else:
                    col_name = 'specifications'
                    data = dict()
                    specs = ['log_elems', 'memory', 'pll', 'multiplier', 'pins']
                    for i in range(4, 9):
                        if self.ui.tableWidget.item(current_row, i) is not None:
                            if self.ui.tableWidget.item(current_row, i).text().isdigit():
                                data[specs[i - 4]] = int(self.ui.tableWidget.item(current_row, i).text())
                            else:
                                self.hardware_table()
                                return QMessageBox.critical(self, 'Ошибка',
                                                            'При изменения данных произошла ошибка.\nДопустим ввод только чисел.',
                                                            QMessageBox.Ok)

                    hardware_response = helpers.patch_request('hardware', col_name, data, update_id)
                if hardware_response['detail'] == "OK":
                    QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
                else:
                    QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка', QMessageBox.Ok)
                    self.hardware_table()
            elif self.ui.tableWidget.columnCount() == 12:
                reqs = helpers.get_request('request')
                current_header = self.ui.tableWidget.horizontalHeaderItem(current_col).text()
                for x in reqs:
                    if x['id'] == update_id:
                        dat = x
                dat[current_header] = item
                data_json = json.dumps(dat, ensure_ascii=False)
                response = requests.patch(f"{DB_URL}/request/{update_id}",
                                          headers={
                                              'Authorization': DB_ACCESS_TOKEN,
                                          },
                                          data=data_json
                                          ).json()
                if response['detail'] == "OK":
                    QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
                else:
                    QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка}', QMessageBox.Ok)
                    self.request_table()

    def request_table(self):
        """
        Загрузка таблицы запросов
        """
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(12)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["ID", "status", "location", "taken_date", "issued_by", "comment", "created", "user", "return_date",
             "hardware", "stock", "count"])

        reqs = helpers.get_request('request?joined=True')

        self.ui.tableWidget.setRowCount(len(reqs))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.status_box = QComboBox()
            statuses = ['new', 'taken', 'completed', 'canceled']
            self.status_box.addItems(statuses)
            self.boxes[row] = self.status_box
            current_status = reqs[row]["status"]
            index = statuses.index(current_status)
            self.status_box.setCurrentIndex(index)
            self.status_box.currentIndexChanged.connect(lambda _, row=row: self.patch_status(row))
            self.ui.tableWidget.setCellWidget(row, 1, self.status_box)
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(str(reqs[row]["location"])))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(reqs[row]["taken_date"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(str(reqs[row]["issued_by"])))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(reqs[row]["comment"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(reqs[row]["created"]))
            self.ui.tableWidget.item(row, 6).setFlags(QtCore.Qt.ItemIsEnabled)
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(str(reqs[row]["user"])))
            self.ui.tableWidget.setItem(row, 8, QTableWidgetItem(reqs[row]["return_date"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(reqs[row]["id"])))
            self.ui.tableWidget.setItem(row, 9, QTableWidgetItem(str(reqs[row]["hardware"])))
            self.ui.tableWidget.setItem(row, 10, QTableWidgetItem(str(reqs[row]["stock"])))
            self.ui.tableWidget.setItem(row, 11, QTableWidgetItem(str(reqs[row]["count"])))

    def patch_status(self, tableRow):
        """
        Отправка изменений в базу данных

        Args:
            tableRow: измененный ряд
        """
        item = self.boxes[tableRow].currentText()
        update_id = int(self.ui.tableWidget.item(tableRow, 0).text())
        reqs = helpers.get_request('request')
        for x in reqs:
            if x['id'] == update_id:
                dat = x
        dat['status'] = item
        data_json = json.dumps(dat, ensure_ascii=False)
        response = requests.patch(f"{DB_URL}/request/{update_id}",
                                  headers={
                                      'Authorization': DB_ACCESS_TOKEN,
                                  },
                                  data=data_json
                                  ).json()
        if response['detail'] == "OK":
            QMessageBox.information(self, 'Успех', 'Данные были успешно изменены')
        else:
            QMessageBox.critical(self, 'Ошибка', 'При изменения данных произошла ошибка}', QMessageBox.Ok)
            self.request_table()

    def switch_theme(self):
        """
        Переключение темы приложения.
        """
        if self.styleSheet() != "":
            self.setStyleSheet("")
        else:
            self.setStyleSheet(qdarkstyle.load_stylesheet_pyqt5())

    def add_hardware(self):
        """
        Открытие окна добавления нового оборудования.
        """
        self.hard_dialog = HardDialog()
        self.hard_dialog.show()

    def add_user(self):
        """
        Открытие окна добавления нового пользователя.
        """
        self.user_dialog = UserDialog()
        self.user_dialog.ui.surname_line.setText(self.ui.search_line.text())
        self.ui.search_line.clear()
        self.user_dialog.show()

    def search_user(self):
        """
        Загрузка таблицы пользователей и поиск по фамилии.
        """
        self.user_table()
        last_name = self.ui.search_line.text().lower()
        in_table = False
        for row in range(self.ui.tableWidget.rowCount()):
            if last_name == self.ui.tableWidget.item(row, 2).text().lower():
                in_table = True
            else:
                self.ui.tableWidget.setRowHidden(row, True)
        if not in_table:
            self.user_table()
            button = QMessageBox.question(
                self,
                'Ошибка',
                'Данного пользователя не существует. Хотите добавить?',
                QMessageBox.StandardButton.Yes |
                QMessageBox.StandardButton.No
            )
            if button == QMessageBox.StandardButton.Yes:
                self.add_user()
        else:
            self.ui.search_line.clear()

    def delete(self):
        """
        Удаление записи в базе данных.
        """
        current_row = self.ui.tableWidget.currentRow()
        if current_row < 0:
            return QMessageBox.warning(self, 'Внимание', 'Пожалуйста выберете запись для удаления')

        button = QMessageBox.question(
            self,
            'Подтверждение',
            'Вы уверены, что хотите удалить данную запись?',
            QMessageBox.StandardButton.Yes |
            QMessageBox.StandardButton.No
        )
        if button == QMessageBox.StandardButton.Yes:
            delete_id = int(self.ui.tableWidget.item(current_row, 0).text())
            if self.ui.tableWidget.columnCount() == 8:
                user_response = requests.delete(f"{DB_URL}/user/{delete_id}",
                                                headers={
                                                    'Authorization': DB_ACCESS_TOKEN}
                                                ).json()
                if 'detail' in user_response:
                    return QMessageBox.warning(self, 'Ошибка',
                                               'Пользователя нельзя удалить, так как у него есть запросы')
                self.ui.tableWidget.removeRow(current_row)
            if self.ui.tableWidget.columnCount() == 9:
                hardware_response = requests.delete(f"{DB_URL}/hardware/{delete_id}",
                                                    headers={
                                                        'Authorization': DB_ACCESS_TOKEN}
                                                    ).json()
                if 'detail' in hardware_response:
                    return QMessageBox.warning(self, 'Ошибка', 'Плату удалить нельзя, так как она находится на складе')
                self.ui.tableWidget.removeRow(current_row)
            if self.ui.tableWidget.columnCount() == 12:
                return QMessageBox.warning(self, 'Ошибка', 'Данная функция пока невозможна')

    def user_table(self):
        """
        загрузка таблицы пользователей
        """
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(8)
        self.ui.tableWidget.setHorizontalHeaderLabels(
            ["ID", "Имя", "Фамилия", "Отчество", "Уровень доступа", "Телефон", "Почта", "Комментарий"])

        users = helpers.get_request('user')

        self.ui.tableWidget.setRowCount(len(users))
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(users[row]["first_name"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(users[row]["last_name"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(users[row]["patronymic"]))
            self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(users[row]["comment"]))
            self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(users[row]["type"]))
            self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(users[row]["phone"]))
            self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(users[row]["email"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(users[row]["id"])))

    def hardware_table(self):
        """
        загрузка таблицы оборудования
        """
        self.ui.tableWidget.clearContents()
        self.ui.tableWidget.setColumnCount(9)
        self.ui.tableWidget.setHorizontalHeaderLabels(["ID", "Тип", "Название", "Описание", "log_elems",
                                                       "memory", "pll", "multiplier", "pins"
                                                       ])
        hardware = helpers.get_request('hardware')

        self.ui.tableWidget.setRowCount(len(hardware))
        self.ui.hardware_list.clear()
        for row in range(self.ui.tableWidget.rowCount()):
            self.ui.tableWidget.setRowHidden(row, False)
            self.ui.tableWidget.setItem(row, 1, QTableWidgetItem(hardware[row]["type"]))
            self.ui.tableWidget.setItem(row, 2, QTableWidgetItem(hardware[row]["name"]))
            self.ui.tableWidget.setItem(row, 3, QTableWidgetItem(hardware[row]["description"]))
            self.ui.tableWidget.setItem(row, 0, QTableWidgetItem(str(hardware[row]['id'])))

            specs = hardware[row]["specifications"]
            if specs:
                self.ui.tableWidget.setItem(row, 4, QTableWidgetItem(str(specs.get('log_elems'))))
                self.ui.tableWidget.setItem(row, 5, QTableWidgetItem(str(specs.get('memory'))))
                self.ui.tableWidget.setItem(row, 6, QTableWidgetItem(str(specs.get('pll'))))
                self.ui.tableWidget.setItem(row, 7, QTableWidgetItem(str(specs.get('multiplier'))))
                self.ui.tableWidget.setItem(row, 8, QTableWidgetItem(str(specs.get('pins'))))
            self.ui.hardware_list.addItem(QListWidgetItem(hardware[row]["name"]))

    def save_excel_file(self):
        """
        Создание эксель файла
        """
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".xls(*.xls)")
        wbk = Workbook()
        sheet = wbk.add_sheet("sheet", cell_overwrite_ok=True)
        self.make_sheet(sheet)
        wbk.save(filepath)

    def make_sheet(self, sheet):
        """
        Создание Листа в экселе
        """
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            sheet.write(0, currentColumn, str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text()))
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            for currentRow in range(self.ui.tableWidget.rowCount()):
                if currentColumn == 1 and self.ui.tableWidget.columnCount() == 12:
                    try:
                        teext = self.boxes[currentRow].currentText()
                        sheet.write(currentRow + 1, currentColumn, teext)
                    except AttributeError:
                        pass
                else:
                    try:
                        teext = str(self.ui.tableWidget.item(currentRow, currentColumn).text())
                        sheet.write(currentRow + 1, currentColumn, teext)
                    except AttributeError:
                        pass

    def save_doc_file(self):
        """
        Создание док файла
        """
        filepath, _ = QFileDialog.getSaveFileName(self, 'Save File', '', ".docx(*.docx)")
        doc = Document()
        table = doc.add_table(rows=self.ui.tableWidget.rowCount() + 1, cols=self.ui.tableWidget.columnCount())
        table.style = 'Table Grid'
        for currentColumn in range(self.ui.tableWidget.columnCount()):
            cell = table.cell(0, currentColumn)
            cell.text = str(self.ui.tableWidget.horizontalHeaderItem(currentColumn).text())
        for row in range(self.ui.tableWidget.rowCount()):
            for col in range(self.ui.tableWidget.columnCount()):
                cell = table.cell(row + 1, col)
                if col == 1 and self.ui.tableWidget.columnCount() == 12:
                    try:
                        teext = self.boxes[row].currentText()
                        cell.text = teext
                    except AttributeError:
                        pass
                else:
                    try:
                        cell.text = str(self.ui.tableWidget.item(row, col).text())
                    except AttributeError:
                        pass
        doc.save(filepath)


if __name__ == '__main__':
    app = QtWidgets.QApplication(sys.argv)
    window = LoginWindow()
    window.show()
    sys.exit(app.exec_())
